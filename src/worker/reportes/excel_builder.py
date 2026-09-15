"""Generacion del Excel de reporte (dos hojas: Impresiones/Clicks).

Usa `openpyxl` en modo `write_only` para escritura incremental (streaming), evitando
cargar todo el reporte en memoria (FR-009 / FR-011).
"""

from __future__ import annotations

import io
from typing import Iterable

import openpyxl

_NA = "N/D"

_ENCABEZADOS = ("timestamp", "usuario_id")


def _valor_o_nd(valor):
    if valor is None:
        return _NA
    if hasattr(valor, "isoformat"):
        # Excel no soporta timezone-aware datetimes; se serializa como ISO-8601 string.
        return valor.isoformat()
    return valor


def construir_reporte(eventos: Iterable) -> io.BytesIO:
    """Construye el workbook de dos hojas a partir de un iterable de eventos.

    Cada evento debe exponer `tipo_evento`, `timestamp`, `usuario_id`.
    """

    wb = openpyxl.Workbook(write_only=True)

    hoja_impresiones = wb.create_sheet("Impresiones")
    hoja_clicks = wb.create_sheet("Clicks")

    hoja_impresiones.append(_ENCABEZADOS)
    hoja_clicks.append(_ENCABEZADOS)

    for evento in eventos:
        fila = (_valor_o_nd(evento.timestamp), _valor_o_nd(evento.usuario_id))
        if evento.tipo_evento == "impresion":
            hoja_impresiones.append(fila)
        elif evento.tipo_evento == "click":
            hoja_clicks.append(fila)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
