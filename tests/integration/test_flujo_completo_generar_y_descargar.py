"""Integration test: flujo completo e2e uniendo las tres features.

reporte.generar (RabbitMQ) -> Consumer real (repository + excel_builder + persistencia
MinIO + autorizacion_repository) -> reporte.listo (RabbitMQ) -> descarga autorizada del
Excel resultante via Nginx + auth-service + MinIO (JWT real), tal como lo haria un
usuario final.

Requiere la infraestructura real de docker-compose.test.yml: rabbitmq-test,
postgres-reportes-test, minio-test, auth-service-test, nginx-test.
"""

from __future__ import annotations

import json
import os
import time
import uuid
from datetime import datetime, timedelta, timezone

import httpx
import jwt
import pika
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

from src.models.anuncio import Base
from src.worker.events.consumer import Consumer
from src.worker.events.publisher import Publisher
from src.worker.reportes import autorizacion_repository as autorizacion_repository_module
from src.worker.reportes import repository as repository_module
from src.worker.reportes.excel_builder import construir_reporte
from src.worker.storage import persistencia as persistencia_module
from src.worker.storage.minio_client import (
    aplicar_politica_lectura_interna,
    asegurar_bucket,
    crear_cliente_minio,
)

RABBITMQ_URL = os.environ.get("TEST_RABBITMQ_URL", "amqp://guest:guest@localhost:5673/")
DATABASE_URL = os.environ.get(
    "TEST_DATABASE_URL", "postgresql+psycopg://reportes:reportes@localhost:5433/reportes"
)
MINIO_ENDPOINT = os.environ.get("TEST_MINIO_ENDPOINT", "localhost:9002")
MINIO_ACCESS_KEY = os.environ.get("TEST_MINIO_ACCESS_KEY", "reportes-test")
MINIO_SECRET_KEY = os.environ.get("TEST_MINIO_SECRET_KEY", "reportes-test-secret")
MINIO_BUCKET = os.environ.get("TEST_MINIO_BUCKET", "reportes-test")
NGINX_BASE_URL = os.environ.get("TEST_NGINX_URL", "http://localhost:8080")
JWT_SECRET = os.environ.get("TEST_JWT_SECRET", "testing-secret")

QUEUE_GENERAR = "test.flujo_completo.reporte.generar"
QUEUE_LISTO = "test.flujo_completo.reporte.listo"


def _token(sub: str, exp_delta: int = 3600) -> str:
    return jwt.encode(
        {"sub": sub, "exp": int(time.time()) + exp_delta}, JWT_SECRET, algorithm="HS256"
    )


class _RepositoryAdapter:
    def __init__(self, session: Session):
        self._session = session

    def obtener_eventos_anuncio(self, anuncio_id, fecha_desde, fecha_hasta):
        return repository_module.obtener_eventos_anuncio(
            self._session, anuncio_id, fecha_desde, fecha_hasta
        )

    def anuncio_existe(self, anuncio_id):
        return repository_module.anuncio_existe(self._session, anuncio_id)


class _PersistenciaAdapter:
    def __init__(self, client, bucket: str):
        self._client = client
        self._bucket = bucket

    def subir_reporte(self, anuncio_id, solicitud_id, archivo):
        return persistencia_module.subir_reporte(
            anuncio_id, solicitud_id, archivo, client=self._client, bucket=self._bucket
        )


class _AutorizacionRepositoryAdapter:
    def __init__(self, session: Session):
        self._session = session

    def registrar_autorizacion(self, *, solicitud_id, anuncio_id, usuario_solicitante, bucket, key):
        autorizacion_repository_module.registrar_autorizacion(
            self._session,
            solicitud_id=solicitud_id,
            anuncio_id=anuncio_id,
            usuario_solicitante=usuario_solicitante,
            bucket=bucket,
            key=key,
        )


class _ExcelBuilderAdapter:
    @staticmethod
    def construir_reporte(eventos):
        return construir_reporte(eventos)


@pytest.fixture(scope="module")
def db_engine():
    engine = create_engine(DATABASE_URL)
    Base.metadata.create_all(engine)
    return engine


@pytest.fixture(scope="module")
def minio_cliente():
    client = crear_cliente_minio(
        endpoint=MINIO_ENDPOINT,
        access_key=MINIO_ACCESS_KEY,
        secret_key=MINIO_SECRET_KEY,
        secure=False,
    )
    asegurar_bucket(client, MINIO_BUCKET)
    aplicar_politica_lectura_interna(client, MINIO_BUCKET)
    return client


@pytest.fixture()
def rabbit_channel():
    connection = pika.BlockingConnection(pika.URLParameters(RABBITMQ_URL))
    channel = connection.channel()
    channel.queue_declare(queue=QUEUE_GENERAR, durable=True)
    channel.queue_declare(queue=QUEUE_LISTO, durable=True)
    channel.queue_purge(QUEUE_GENERAR)
    channel.queue_purge(QUEUE_LISTO)
    yield channel
    channel.queue_purge(QUEUE_GENERAR)
    channel.queue_purge(QUEUE_LISTO)
    connection.close()


def _insertar_evento(engine, anuncio_id, tipo_evento, timestamp, usuario_id="user-1"):
    with engine.begin() as conn:
        conn.execute(
            text(
                "INSERT INTO anuncio (id, anuncio_id, tipo_evento, timestamp, usuario_id) "
                "VALUES (:id, :anuncio_id, :tipo_evento, :timestamp, :usuario_id)"
            ),
            {
                "id": str(uuid.uuid4()),
                "anuncio_id": anuncio_id,
                "tipo_evento": tipo_evento,
                "timestamp": timestamp,
                "usuario_id": usuario_id,
            },
        )


def _publicar_solicitud(channel, payload: dict) -> None:
    channel.basic_publish(
        exchange="",
        routing_key=QUEUE_GENERAR,
        body=json.dumps(payload).encode("utf-8"),
        properties=pika.BasicProperties(content_type="application/json", delivery_mode=2),
    )


def _consumir_un_mensaje(engine, minio_cliente, channel, *, timeout_s: float = 10.0):
    """Levanta el Consumer real (con persistencia + autorizacion_repository, igual que
    src/worker/main.py) y consume exactamente un mensaje de QUEUE_GENERAR."""

    with Session(engine) as session:
        repository = _RepositoryAdapter(session)
        publisher = Publisher(channel, QUEUE_LISTO)
        persistencia = _PersistenciaAdapter(minio_cliente, MINIO_BUCKET)
        autorizacion_repository = _AutorizacionRepositoryAdapter(session)
        consumer = Consumer(
            channel=channel,
            repository=repository,
            excel_builder=_ExcelBuilderAdapter,
            publisher=publisher,
            persistencia=persistencia,
            autorizacion_repository=autorizacion_repository,
        )
        channel.basic_qos(prefetch_count=1)
        channel.basic_consume(queue=QUEUE_GENERAR, on_message_callback=consumer.procesar_mensaje)

        deadline = time.monotonic() + timeout_s
        while time.monotonic() < deadline:
            channel.connection.process_data_events(time_limit=0.5)
        channel.stop_consuming()


def _leer_mensaje_listo(channel) -> dict | None:
    method, properties, body = channel.basic_get(queue=QUEUE_LISTO, auto_ack=True)
    if body is None:
        return None
    return json.loads(body)


def test_flujo_completo_generar_persistir_autorizar_y_descargar(
    db_engine, minio_cliente, rabbit_channel
):
    """Flujo real de punta a punta: publica reporte.generar, el Consumer real
    consulta eventos, genera el Excel, lo sube a MinIO, registra la autorizacion,
    publica reporte.listo, y finalmente el usuario autorizado descarga el reporte
    via Nginx + auth-service usando un JWT, obteniendo el Excel real generado."""

    anuncio_id = f"anuncio-{uuid.uuid4()}"
    solicitud_id = str(uuid.uuid4())
    usuario = "usuario-flujo-completo"
    ahora = datetime.now(timezone.utc)
    _insertar_evento(db_engine, anuncio_id, "impresion", ahora - timedelta(days=1), usuario)
    _insertar_evento(db_engine, anuncio_id, "click", ahora - timedelta(hours=12), usuario)

    payload = {
        "solicitud_id": solicitud_id,
        "anuncio_id": anuncio_id,
        "fecha_desde": (ahora - timedelta(days=2)).isoformat(),
        "fecha_hasta": (ahora + timedelta(days=1)).isoformat(),
        "usuario_solicitante": usuario,
    }
    _publicar_solicitud(rabbit_channel, payload)
    _consumir_un_mensaje(db_engine, minio_cliente, rabbit_channel)

    # 1. reporte.listo fue publicado con estado "generado" y referencia al archivo real.
    resultado = _leer_mensaje_listo(rabbit_channel)
    assert resultado is not None
    assert resultado["solicitud_id"] == solicitud_id
    assert resultado["estado"] == "generado"
    assert resultado["referencia_archivo"] == f"{MINIO_BUCKET}/reportes/{anuncio_id}/{solicitud_id}.xlsx"

    # 2. El usuario autorizado descarga el reporte real via Nginx + auth-service, con un
    #    JWT emitido para el mismo usuario_solicitante que la solicitud original.
    resp = httpx.get(
        f"{NGINX_BASE_URL}/reportes/{solicitud_id}",
        headers={"Authorization": f"Bearer {_token(usuario)}"},
    )
    assert resp.status_code == 200

    # 3. El contenido descargado es un Excel valido con los eventos consultados.
    from io import BytesIO

    workbook = load_workbook(BytesIO(resp.content))
    hoja = workbook.active
    filas = list(hoja.iter_rows(values_only=True))
    assert len(filas) >= 2  # al menos encabezado + 1 fila de datos

    # 4. Un usuario distinto no puede descargar el mismo reporte (control de acceso
    #    de 003 aplicado sobre datos reales generados por 001+002).
    resp_no_autorizado = httpx.get(
        f"{NGINX_BASE_URL}/reportes/{solicitud_id}",
        headers={"Authorization": f"Bearer {_token('otro-usuario')}"},
    )
    assert resp_no_autorizado.status_code == 403


def test_flujo_completo_sin_eventos_publica_vacio_y_sigue_siendo_descargable(
    db_engine, minio_cliente, rabbit_channel
):
    """Aun con estado 'vacio' (sin eventos en el rango), el Excel generado (solo con
    encabezados) se persiste, autoriza y puede descargarse normalmente."""

    anuncio_id = f"anuncio-{uuid.uuid4()}"
    solicitud_id = str(uuid.uuid4())
    usuario = "usuario-flujo-vacio"
    ahora = datetime.now(timezone.utc)
    # El anuncio existe pero el evento cae fuera del rango solicitado.
    _insertar_evento(db_engine, anuncio_id, "impresion", ahora - timedelta(days=365), usuario)

    payload = {
        "solicitud_id": solicitud_id,
        "anuncio_id": anuncio_id,
        "fecha_desde": (ahora - timedelta(days=2)).isoformat(),
        "fecha_hasta": (ahora + timedelta(days=1)).isoformat(),
        "usuario_solicitante": usuario,
    }
    _publicar_solicitud(rabbit_channel, payload)
    _consumir_un_mensaje(db_engine, minio_cliente, rabbit_channel)

    resultado = _leer_mensaje_listo(rabbit_channel)
    assert resultado is not None
    assert resultado["estado"] == "vacio"

    resp = httpx.get(
        f"{NGINX_BASE_URL}/reportes/{solicitud_id}",
        headers={"Authorization": f"Bearer {_token(usuario)}"},
    )
    assert resp.status_code == 200
