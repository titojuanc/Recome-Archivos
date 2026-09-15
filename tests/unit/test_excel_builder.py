"""Unit tests: excel_builder.construir_reporte().

Deben FALLAR hasta que exista src/worker/reportes/excel_builder.py (TDD estricto).
"""

from __future__ import annotations

import io
from datetime import datetime, timezone

import openpyxl
import pytest


class EventoFake:
    """Stub minimo de un registro de `anuncio` para no depender de SQLAlchemy aca."""

    def __init__(self, tipo_evento, timestamp, usuario_id):
        self.tipo_evento = tipo_evento
        self.timestamp = timestamp
        self.usuario_id = usuario_id


@pytest.fixture()
def eventos_mixtos():
    ahora = datetime(2026, 1, 1, tzinfo=timezone.utc)
    return [
        EventoFake("impresion", ahora, "user-1"),
        EventoFake("click", ahora, "user-1"),
        EventoFake("impresion", None, None),  # corrupto -> "N/D"
    ]


def _leer_workbook(stream: io.BytesIO):
    stream.seek(0)
    return openpyxl.load_workbook(stream)


def test_genera_workbook_con_hojas_impresiones_y_clicks(eventos_mixtos):
    from src.worker.reportes.excel_builder import construir_reporte

    buffer = construir_reporte(eventos_mixtos)
    wb = _leer_workbook(buffer)

    assert wb.sheetnames == ["Impresiones", "Clicks"]


def test_marca_campos_corruptos_como_nd(eventos_mixtos):
    from src.worker.reportes.excel_builder import construir_reporte

    buffer = construir_reporte(eventos_mixtos)
    wb = _leer_workbook(buffer)

    hoja_impresiones = wb["Impresiones"]
    filas = list(hoja_impresiones.iter_rows(values_only=True))

    # header + 2 filas de impresiones (una valida, una corrupta)
    assert len(filas) == 3
    fila_corrupta = filas[2]
    assert "N/D" in fila_corrupta


def test_genera_hojas_con_solo_encabezados_cuando_no_hay_datos():
    from src.worker.reportes.excel_builder import construir_reporte

    buffer = construir_reporte([])
    wb = _leer_workbook(buffer)

    for nombre_hoja in ("Impresiones", "Clicks"):
        filas = list(wb[nombre_hoja].iter_rows(values_only=True))
        assert len(filas) == 1  # solo encabezado
